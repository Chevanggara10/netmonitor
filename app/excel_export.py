"""
Export riwayat check ke file Excel (.xlsx) -- dipakai untuk download
dari dashboard, dan sebagai dataset input untuk training model forecast
(training/train_forecast.py).

Dua sheet:
- "Raw Checks": data mentah, sama isinya dengan export CSV (reporting.py).
- "Hourly Aggregate": agregat per jam (reuse get_hourly_aggregate),
  supaya training script punya fitur time-series siap pakai tanpa
  perlu resample ulang dari jutaan baris mentah.
"""
import io
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.models import CheckResult
from app.reporting import get_hourly_aggregate

RAW_HEADER = [
    "target_name", "checked_at", "status_code", "response_time_ms",
    "is_up", "error_message", "is_anomaly", "anomaly_z_score",
]
HOURLY_HEADER = ["hour", "avg_response_time_ms", "total_checks", "uptime_percent"]

# Batas baris mentah per export (yang terbaru). Interval 5 detik selama
# setahun = jutaan baris: tanpa batas, satu request bisa menghabiskan memori
# server dan menggantung dashboard. Data yang dipotong ditandai di sheet "Info".
MAX_EXPORT_ROWS = 200_000
MAX_EXPORT_DAYS = 3650

_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_cell(value):
    """
    Nilai teks berasal dari luar (nama target dari user, pesan error dari
    server target). Tanpa sanitasi: (1) teks berawalan '=' ditulis sebagai
    RUMUS dan dieksekusi Excel saat dibuka (formula injection), (2) karakter
    kontrol membuat openpyxl melempar IllegalCharacterError -> export 500.
    """
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


async def export_checks_to_excel(target_id: int, target_name: str, db: AsyncSession, days: int = 30) -> bytes:
    days = max(1, min(int(days), MAX_EXPORT_DAYS))
    since = datetime.utcnow() - timedelta(days=days)

    # Ambil MAX+1 baris TERBARU untuk mendeteksi pemotongan, lalu balik ke
    # urutan kronologis (training butuh urutan waktu naik).
    result = await db.execute(
        select(CheckResult)
        .where(CheckResult.target_id == target_id, CheckResult.checked_at >= since)
        .order_by(CheckResult.checked_at.desc())
        .limit(MAX_EXPORT_ROWS + 1)
    )
    checks = list(result.scalars().all())
    truncated = len(checks) > MAX_EXPORT_ROWS
    checks = list(reversed(checks[:MAX_EXPORT_ROWS]))

    wb = Workbook()
    raw_ws = wb.active
    raw_ws.title = "Raw Checks"
    raw_ws.append(RAW_HEADER)
    safe_name = _safe_cell(target_name)
    for c in checks:
        raw_ws.append([
            safe_name, c.checked_at.isoformat(), c.status_code, c.response_time_ms,
            c.is_up, _safe_cell(c.error_message or ""), c.is_anomaly, c.anomaly_z_score,
        ])

    hourly_ws = wb.create_sheet("Hourly Aggregate")
    hourly_ws.append(HOURLY_HEADER)
    hourly_rows = await get_hourly_aggregate(target_id, db, days=days)
    for row in hourly_rows:
        hourly_ws.append([row["hour"], row["avg_response_time_ms"], row["total_checks"], row["uptime_percent"]])

    if truncated:
        info_ws = wb.create_sheet("Info")
        info_ws.append(["catatan"])
        info_ws.append([
            f"Data mentah dipotong ke {MAX_EXPORT_ROWS} baris TERBARU dari rentang {days} hari. "
            "Kecilkan parameter days untuk mendapatkan seluruh data."
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
