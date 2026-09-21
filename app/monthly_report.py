"""
Laporan bulanan otomatis ke Telegram.

Alur: tiap jam scheduler memanggil ensure_monthly_reports(). Mulai tanggal 1
pukul REPORT_SEND_HOUR (zona REPORT_TIMEZONE, default Asia/Jakarta), setiap
user yang sudah menghubungkan Telegram menerima laporan BULAN LALU:
ringkasan evaluasi (analis AI) + grafik PNG + file Excel.

Ketahanan:
- Idempoten lewat tabel monthly_report_log (unik user+periode): restart /
  job ganda tidak mengirim dua kali.
- Catch-up: kalau server mati pada tanggal 1, laporan dikirim begitu server
  hidup lagi (kapan pun di bulan berjalan).
- Pengiriman gagal dicoba lagi tiap jam sampai berhasil.
- Satu user bermasalah tidak menghalangi user lain; tidak ada exception
  yang lolos ke scheduler.
"""
import asyncio
import io
import logging
import os
import threading
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")  # tanpa layar; harus sebelum pyplot
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.ai_analyst import analyze, render_evaluation_text  # noqa: E402
from app.database import AsyncSessionLocal  # noqa: E402
from app.excel_export import _safe_cell  # noqa: E402
from app.ml_forecast import get_forecast  # noqa: E402
from app.models import CheckResult, MonitorTarget, MonthlyReportLog, User  # noqa: E402
from app.telegram_bot import send_alert, send_document, send_photos  # noqa: E402

logger = logging.getLogger("netmonitor.monthly_report")

MONTH_NAMES = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus",
               "September", "Oktober", "November", "Desember"]
MAX_CHART_TARGETS = 5
SEND_TIMEOUT_SECONDS = 120

_render_lock = threading.Lock()  # matplotlib tidak thread-safe
_ensure_lock = asyncio.Lock()


@dataclass
class ReportResult:
    status: str  # sent | already_sent | not_linked | inactive | no_targets | no_data | failed
    detail: str = ""


# ---------- periode & zona waktu ----------

def report_timezone():
    name = os.environ.get("REPORT_TIMEZONE", "Asia/Jakarta")
    try:
        return ZoneInfo(name)
    except Exception:
        logger.warning(f"Zona waktu '{name}' tidak dikenal, memakai UTC+7 (WIB).")
        return timezone(timedelta(hours=7), "WIB")


def _send_hour() -> int:
    try:
        return max(0, min(23, int(os.environ.get("REPORT_SEND_HOUR", "8"))))
    except ValueError:
        return 8


def previous_period(now_local: datetime) -> str:
    last_day_prev = now_local.replace(day=1) - timedelta(days=1)
    return f"{last_day_prev.year:04d}-{last_day_prev.month:02d}"


def _to_utc_naive(dt: datetime) -> datetime:
    return dt.astimezone(timezone.utc).replace(tzinfo=None)


def period_bounds_utc(period: str, tz) -> tuple[datetime, datetime]:
    """[awal, akhir) periode 'YYYY-MM' dalam zona laporan, dikonversi ke UTC naif (format DB)."""
    year, month = (int(x) for x in period.split("-"))
    ny, nm = (year + 1, 1) if month == 12 else (year, month + 1)
    return _to_utc_naive(datetime(year, month, 1, tzinfo=tz)), _to_utc_naive(datetime(ny, nm, 1, tzinfo=tz))


def period_label(period: str) -> str:
    year, month = (int(x) for x in period.split("-"))
    return f"{MONTH_NAMES[month - 1]} {year}"


def _prev_period_of(period: str) -> str:
    year, month = (int(x) for x in period.split("-"))
    return f"{year - 1:04d}-12" if month == 1 else f"{year:04d}-{month - 1:02d}"


def _num(value):
    return None if value is None or pd.isna(value) else float(value)


# ---------- statistik ----------

async def _load_frame(db, target_id: int, start: datetime, end: datetime) -> pd.DataFrame:
    rows = (await db.execute(
        select(CheckResult.checked_at, CheckResult.is_up, CheckResult.response_time_ms, CheckResult.is_anomaly)
        .where(CheckResult.target_id == target_id, CheckResult.checked_at >= start, CheckResult.checked_at < end)
        .order_by(CheckResult.checked_at.asc())
    )).all()
    df = pd.DataFrame(rows, columns=["checked_at", "is_up", "rt", "is_anomaly"])
    if not df.empty:
        df["checked_at"] = pd.to_datetime(df["checked_at"])
        df["is_up"] = df["is_up"].astype(bool)
        df["is_anomaly"] = df["is_anomaly"].astype(bool)
        df["rt"] = pd.to_numeric(df["rt"], errors="coerce")
        df["rt_up"] = df["rt"].where(df["is_up"])
    return df


def _basic(df: pd.DataFrame) -> tuple[float | None, float | None]:
    if df.empty:
        return None, None
    return df["is_up"].mean() * 100, _num(df["rt_up"].mean())


async def collect_target_stats(db, target, start: datetime, end: datetime,
                               prev_start: datetime, prev_end: datetime, tz) -> dict:
    stats = {
        "name": target.name, "url": target.url, "check_type": target.check_type,
        "total_checks": 0, "up_checks": 0, "uptime_pct": None, "avg_rt": None, "p95_rt": None, "max_rt": None,
        "incident_count": 0, "downtime_minutes": 0.0, "longest_incident_minutes": 0.0, "anomaly_count": 0,
        "prev_uptime_pct": None, "prev_avg_rt": None, "busiest_hour": None,
        "forecast_avg_ms": None, "forecast_series": [], "daily": [], "incidents": [],
    }

    prev_df = await _load_frame(db, target.id, prev_start, prev_end)
    stats["prev_uptime_pct"], stats["prev_avg_rt"] = _basic(prev_df)

    df = await _load_frame(db, target.id, start, end)
    if df.empty:
        return stats

    df["local"] = df["checked_at"].dt.tz_localize("UTC").dt.tz_convert(tz)
    stats["total_checks"] = int(len(df))
    stats["up_checks"] = int(df["is_up"].sum())
    stats["uptime_pct"] = stats["up_checks"] / stats["total_checks"] * 100
    stats["avg_rt"] = _num(df["rt_up"].mean())
    stats["p95_rt"] = _num(df["rt_up"].quantile(0.95))
    stats["max_rt"] = _num(df["rt_up"].max())
    stats["anomaly_count"] = int(df["is_anomaly"].sum())

    hourly = df.groupby(df["local"].dt.hour)["rt_up"].mean().dropna()
    if not hourly.empty:
        stats["busiest_hour"] = int(hourly.idxmax())

    # insiden = rangkaian check DOWN berurutan
    interval = df["checked_at"].diff().median()
    if pd.isna(interval) or interval <= pd.Timedelta(0):
        interval = pd.Timedelta(seconds=60)
    down = ~df["is_up"]
    run_id = (down != down.shift()).cumsum()
    period_end = pd.Timestamp(end)
    for _, group in df[down].groupby(run_id[down]):
        first, last = group.index[0], group.index[-1]
        began = df.at[first, "checked_at"]
        ended = df.at[last + 1, "checked_at"] if last + 1 < len(df) else min(df.at[last, "checked_at"] + interval, period_end)
        minutes = (ended - began).total_seconds() / 60
        stats["incidents"].append({
            "start": df.at[first, "local"].isoformat(timespec="minutes"),
            "end": ended.tz_localize("UTC").tz_convert(tz).isoformat(timespec="minutes"),
            "minutes": minutes,
        })
    stats["incident_count"] = len(stats["incidents"])
    stats["downtime_minutes"] = float(sum(i["minutes"] for i in stats["incidents"]))
    stats["longest_incident_minutes"] = float(max((i["minutes"] for i in stats["incidents"]), default=0.0))

    for day, g in df.groupby(df["local"].dt.date):
        stats["daily"].append({
            "date": day.isoformat(), "checks": int(len(g)), "uptime_pct": float(g["is_up"].mean() * 100),
            "avg_rt": _num(g["rt_up"].mean()), "p95_rt": _num(g["rt_up"].quantile(0.95)),
            "anomalies": int(g["is_anomaly"].sum()),
        })

    try:
        forecast = await get_forecast(target.id, db, horizon_hours=168)
        if forecast["source"] == "trained_model":
            stats["forecast_series"] = list(zip(forecast["timestamps"], forecast["forecast_ms"]))
            stats["forecast_avg_ms"] = float(sum(forecast["forecast_ms"]) / len(forecast["forecast_ms"]))
    except Exception:
        logger.exception(f"Forecast untuk laporan target {target.id} gagal (dilewati)")

    return stats


# ---------- grafik ----------

def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=110, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def _uptime_color(pct: float) -> str:
    return "#2e9e5b" if pct >= 99.9 else "#e0a526" if pct >= 99.0 else "#d64545"


def _chart_fleet(label: str, with_data: list[dict]) -> dict:
    ordered = sorted(with_data, key=lambda t: t["uptime_pct"])
    fig, ax = plt.subplots(figsize=(7.5, 1.2 + 0.55 * len(ordered)))
    ax.barh([t["name"][:28] for t in ordered], [t["uptime_pct"] for t in ordered],
            color=[_uptime_color(t["uptime_pct"]) for t in ordered])
    lo = min(t["uptime_pct"] for t in ordered)
    ax.set_xlim(max(0, min(lo - 1, 98)), 100.3)
    for i, t in enumerate(ordered):
        rt = f" | {t['avg_rt']:.0f} ms" if t["avg_rt"] else ""
        ax.text(min(t["uptime_pct"] + 0.05, 100.2), i, f"{t['uptime_pct']:.2f}%{rt}", va="center", fontsize=8)
    ax.axvline(99.9, color="#888", linestyle="--", linewidth=0.8)
    ax.set_title(f"Ketersediaan per target - {label}")
    ax.set_xlabel("Uptime (%)  (garis putus-putus = 99,9%)")
    return {"name": "ringkasan-target.png", "bytes": _png(fig), "caption": f"Ringkasan ketersediaan seluruh target - {label}"}


def _chart_response(label: str, t: dict, tz_label: str) -> dict:
    days = [pd.Timestamp(d["date"]) for d in t["daily"]]
    fig, ax = plt.subplots(figsize=(7.5, 3.4))
    ax.plot(days, [d["avg_rt"] for d in t["daily"]], color="#2b6cb0", marker="o", markersize=3, label="rata-rata")
    ax.plot(days, [d["p95_rt"] for d in t["daily"]], color="#9aa5b1", linestyle="--", label="p95")
    flagged = [(pd.Timestamp(d["date"]), d["avg_rt"]) for d in t["daily"] if d["anomalies"] > 0 and d["avg_rt"]]
    if flagged:
        ax.scatter(*zip(*flagged), color="#d64545", zorder=5, label="hari ada anomali")
    ax.set_title(f"Response time harian - {t['name']}")
    ax.set_ylabel("ms")
    ax.legend(fontsize=8)
    fig.autofmt_xdate()
    ax.grid(alpha=0.25)
    caption = (f"{t['name']} - respons rata-rata {t['avg_rt']:.0f} ms, p95 {t['p95_rt']:.0f} ms ({label}, zona {tz_label})"
               if t["avg_rt"] else f"{t['name']} - respons harian ({label})")
    return {"name": f"respons-{t['name'][:20]}.png", "bytes": _png(fig), "caption": caption}


def _chart_uptime(label: str, t: dict) -> dict:
    days = [pd.Timestamp(d["date"]) for d in t["daily"]]
    values = [d["uptime_pct"] for d in t["daily"]]
    fig, ax = plt.subplots(figsize=(7.5, 3.2))
    ax.bar(days, values, color=[_uptime_color(v) for v in values], width=0.8)
    ax.set_ylim(max(0, min(values) - 2), 100.5)
    ax.axhline(99.9, color="#888", linestyle="--", linewidth=0.8)
    ax.set_title(f"Uptime harian - {t['name']}")
    ax.set_ylabel("%")
    fig.autofmt_xdate()
    ax.grid(axis="y", alpha=0.25)
    caption = f"{t['name']} - ketersediaan {t['uptime_pct']:.2f}%, {t['incident_count']} insiden ({label})"
    return {"name": f"uptime-{t['name'][:20]}.png", "bytes": _png(fig), "caption": caption}


def _chart_forecast(t: dict) -> dict:
    xs = [pd.Timestamp(ts) for ts, _ in t["forecast_series"]]
    ys = [v for _, v in t["forecast_series"]]
    fig, ax = plt.subplots(figsize=(7.5, 3.2))
    ax.plot(xs, ys, color="#7b5ea7")
    ax.fill_between(xs, ys, alpha=0.12, color="#7b5ea7")
    ax.set_title(f"Prediksi model AI 7 hari ke depan - {t['name']}")
    ax.set_ylabel("ms")
    fig.autofmt_xdate()
    ax.grid(alpha=0.25)
    return {"name": f"prediksi-{t['name'][:20]}.png", "bytes": _png(fig),
            "caption": f"{t['name']} - prediksi response time 7 hari ke depan (rata-rata {sum(ys) / len(ys):.0f} ms)"}


def render_charts(label: str, stats_list: list[dict], tz) -> list[dict]:
    with_data = [t for t in stats_list if t.get("total_checks", 0) > 0]
    if not with_data:
        return []
    tz_label = getattr(tz, "key", None) or str(tz)
    with _render_lock:
        charts = [_chart_fleet(label, with_data)]
        for t in sorted(with_data, key=lambda t: t["uptime_pct"])[:MAX_CHART_TARGETS]:
            charts.append(_chart_response(label, t, tz_label))
            charts.append(_chart_uptime(label, t))
            if t.get("forecast_series"):
                charts.append(_chart_forecast(t))
    return charts


# ---------- Excel ----------

def build_report_workbook(label: str, stats_list: list[dict]) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Ringkasan"
    summary.append([
        "target", "url", "total_check", "uptime_pct", "avg_rt_ms", "p95_rt_ms", "max_rt_ms", "insiden",
        "downtime_menit", "insiden_terpanjang_menit", "anomali", "uptime_bulan_lalu_pct",
        "avg_rt_bulan_lalu_ms", "jam_terlambat", "prediksi_7hari_avg_ms",
    ])
    daily = wb.create_sheet("Harian")
    daily.append(["target", "tanggal", "total_check", "uptime_pct", "avg_rt_ms", "p95_rt_ms", "anomali"])
    incidents = wb.create_sheet("Insiden")
    incidents.append(["target", "mulai", "selesai", "durasi_menit"])

    for t in stats_list:
        name = _safe_cell(t["name"])
        summary.append([
            name, _safe_cell(t["url"]), t["total_checks"], t["uptime_pct"], t["avg_rt"], t["p95_rt"], t["max_rt"],
            t["incident_count"], t["downtime_minutes"], t["longest_incident_minutes"], t["anomaly_count"],
            t["prev_uptime_pct"], t["prev_avg_rt"], t["busiest_hour"], t["forecast_avg_ms"],
        ])
        for d in t["daily"]:
            daily.append([name, d["date"], d["checks"], d["uptime_pct"], d["avg_rt"], d["p95_rt"], d["anomalies"]])
        for i in t["incidents"]:
            incidents.append([name, i["start"], i["end"], i["minutes"]])

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(12, len(str(col[0].value)) + 2))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------- pengiriman & orkestrasi ----------

async def _bounded(coro, default, what: str):
    """Jalankan pengiriman dengan batas waktu; error/timeout -> default (tidak pernah melempar)."""
    try:
        return await asyncio.wait_for(coro, timeout=SEND_TIMEOUT_SECONDS)
    except asyncio.TimeoutError:
        logger.warning(f"Pengiriman {what} timeout")
    except Exception:
        logger.exception(f"Pengiriman {what} gagal")
    return default


async def _get_log(db, user_id: int, period: str) -> MonthlyReportLog | None:
    result = await db.execute(select(MonthlyReportLog).where(
        MonthlyReportLog.user_id == user_id, MonthlyReportLog.period == period))
    return result.scalars().first()


async def _record(db, user_id: int, period: str, status: str, detail: str) -> None:
    row = await _get_log(db, user_id, period)
    if row is None:
        db.add(MonthlyReportLog(user_id=user_id, period=period, status=status, detail=detail[:500]))
    else:
        row.status, row.detail, row.created_at = status, detail[:500], datetime.utcnow()
    await db.commit()


async def generate_and_send_report(db, user: User, *, mode: str = "previous", force: bool = False,
                                   now: datetime | None = None) -> ReportResult:
    """
    mode="previous": laporan bulan lalu (dicatat & idempoten). mode="current": bulan berjalan
    sampai sekarang (untuk uji manual, tidak dicatat). force=True mengirim ulang walau sudah tercatat.
    """
    if not user.is_active:
        return ReportResult("inactive", "Akun nonaktif.")
    if not user.telegram_chat_id:
        return ReportResult("not_linked", "Telegram belum terhubung.")

    tz = report_timezone()
    now_utc = now or datetime.now(timezone.utc).replace(tzinfo=None)
    now_local = now_utc.replace(tzinfo=timezone.utc).astimezone(tz)

    if mode == "previous":
        period = previous_period(now_local)
        start, end = period_bounds_utc(period, tz)
        label = period_label(period)
        if not force:
            existing = await _get_log(db, user.id, period)
            if existing and existing.status in ("sent", "skipped"):
                return ReportResult("already_sent", f"Periode {period} sudah diproses ({existing.status}).")
    else:
        period = f"{now_local.year:04d}-{now_local.month:02d}"
        start, _ = period_bounds_utc(period, tz)
        end = now_utc
        label = f"{period_label(period)} (berjalan)"
    prev_start, prev_end = period_bounds_utc(_prev_period_of(period), tz)

    targets = (await db.execute(
        select(MonitorTarget).where(MonitorTarget.user_id == user.id).order_by(MonitorTarget.id))).scalars().all()
    if not targets:
        return ReportResult("no_targets", "Belum ada target monitoring.")

    stats = [await collect_target_stats(db, t, start, end, prev_start, prev_end, tz) for t in targets]
    if not any(s["total_checks"] > 0 for s in stats):
        if mode == "previous":
            await _record(db, user.id, period, "skipped", "tidak ada data pada periode ini")
        return ReportResult("no_data", f"Tidak ada data monitoring pada {label}.")

    evaluation, narrative = await analyze(label, stats)
    charts = await asyncio.to_thread(render_charts, label, stats, tz)
    workbook = await asyncio.to_thread(build_report_workbook, label, stats)

    text = render_evaluation_text(label, evaluation, narrative)
    if sum(1 for s in stats if s["total_checks"] > 0) > MAX_CHART_TARGETS:
        text += (f"\n\nCatatan: grafik detail hanya untuk {MAX_CHART_TARGETS} target dengan "
                 "ketersediaan terendah; semua target ada di file Excel.")

    chat_id = user.telegram_chat_id
    if not await _bounded(send_alert(chat_id, text), False, "ringkasan laporan"):
        if mode == "previous":
            await _record(db, user.id, period, "failed", "pesan ringkasan gagal terkirim")
        return ReportResult("failed", "Pesan ringkasan gagal terkirim ke Telegram.")

    photos_sent = await _bounded(send_photos(chat_id, [(c["bytes"], c["caption"]) for c in charts]), 0, "grafik")
    doc_ok = await _bounded(
        send_document(chat_id, workbook, f"laporan-{period}.xlsx", f"Data lengkap {label} (Ringkasan, Harian, Insiden)"),
        False, "excel")

    detail = f"foto {photos_sent}/{len(charts)}, excel {'ok' if doc_ok else 'gagal'}"
    if mode == "previous":
        await _record(db, user.id, period, "sent", detail)
    return ReportResult("sent", detail)


async def ensure_monthly_reports(now: datetime | None = None, *, session_factory=AsyncSessionLocal) -> list[tuple[int, str]]:
    """Dipanggil tiap jam: kirim laporan bulan lalu ke semua user terhubung yang belum menerimanya."""
    if _ensure_lock.locked():
        return []
    async with _ensure_lock:
        tz = report_timezone()
        now_utc = now or datetime.now(timezone.utc).replace(tzinfo=None)
        now_local = now_utc.replace(tzinfo=timezone.utc).astimezone(tz)
        if now_local.day == 1 and now_local.hour < _send_hour():
            return []

        results: list[tuple[int, str]] = []
        async with session_factory() as db:
            users = (await db.execute(select(User).where(
                User.is_active == True, User.telegram_chat_id.is_not(None)).order_by(User.id))).scalars().all()  # noqa: E712
            for user in users:
                try:
                    result = await generate_and_send_report(db, user, mode="previous", now=now_utc)
                    results.append((user.id, result.status))
                except Exception:
                    logger.exception(f"Laporan bulanan untuk user {user.id} gagal")
                    results.append((user.id, "error"))
        return results


def schedule_report_jobs(scheduler) -> None:
    scheduler.add_job(ensure_monthly_reports, "cron", minute=5, id="monthly_reports", replace_existing=True,
                      coalesce=True, misfire_grace_time=3600, max_instances=1)
    scheduler.add_job(ensure_monthly_reports, "date", run_date=datetime.now() + timedelta(seconds=120),
                      id="monthly_reports_startup", replace_existing=True)
