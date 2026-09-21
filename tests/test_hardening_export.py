import io
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app.excel_export import export_checks_to_excel
from app.models import CheckResult, MonitorTarget
from tests.conftest import add_check_results


async def test_formula_injection_in_target_name_and_error_message_is_neutralised(db_session, sample_user):
    target = MonitorTarget(user_id=sample_user.id, name='=HYPERLINK("http://evil","x")', url="https://example.com")
    db_session.add(target)
    await db_session.commit()
    await db_session.refresh(target)
    for i, evil in enumerate(["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1+1)"]):
        db_session.add(CheckResult(
            target_id=target.id, checked_at=datetime.utcnow() - timedelta(minutes=i),
            is_up=False, error_message=evil,
        ))
    await db_session.commit()

    content = await export_checks_to_excel(target.id, target.name, db_session, days=30)

    ws = load_workbook(io.BytesIO(content))["Raw Checks"]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != "f", f"formula found in {cell.coordinate}: {cell.value!r}"
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("=", "+", "-", "@")), cell.value


async def test_illegal_control_characters_in_error_message_do_not_break_export(db_session, sample_target):
    db_session.add(CheckResult(
        target_id=sample_target.id, checked_at=datetime.utcnow(),
        is_up=False, error_message="bad\x00byte\x07 and \x1b[31mansi",
    ))
    await db_session.commit()

    content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=30)

    ws = load_workbook(io.BytesIO(content))["Raw Checks"]
    message = ws.cell(row=2, column=6).value
    assert "\x00" not in message and "\x07" not in message
    assert "bad" in message and "ansi" in message


async def test_export_is_capped_to_most_recent_rows_with_info_sheet(db_session, sample_target, monkeypatch):
    monkeypatch.setattr("app.excel_export.MAX_EXPORT_ROWS", 3)
    await add_check_results(db_session, sample_target.id, count=10)

    content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=30)

    wb = load_workbook(io.BytesIO(content))
    raw = wb["Raw Checks"]
    assert raw.max_row == 4  # header + 3 baris terbaru
    checked = [raw.cell(row=r, column=2).value for r in range(2, 5)]
    assert checked == sorted(checked)  # urutan kronologis dipertahankan
    assert "Info" in wb.sheetnames


async def test_export_without_truncation_has_no_info_sheet(db_session, sample_target):
    await add_check_results(db_session, sample_target.id, count=3)

    content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=30)

    assert "Info" not in load_workbook(io.BytesIO(content)).sheetnames


async def test_non_positive_or_huge_days_do_not_crash(db_session, sample_target):
    await add_check_results(db_session, sample_target.id, count=3)

    for days in (0, -5, 10**9):
        content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=days)
        assert "Raw Checks" in load_workbook(io.BytesIO(content)).sheetnames
