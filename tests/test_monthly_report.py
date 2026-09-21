import io
from datetime import datetime, timedelta
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import ml_forecast, monthly_report, telegram_bot
from app.models import CheckResult, MonitorTarget, MonthlyReportLog, User
from app.monthly_report import (
    build_report_workbook, collect_target_stats, ensure_monthly_reports, generate_and_send_report,
    period_bounds_utc, period_label, previous_period, render_charts, report_timezone,
)

TZ = report_timezone()
AUG_START, AUG_END = period_bounds_utc("2026-08", TZ)
JUL_START, JUL_END = period_bounds_utc("2026-07", TZ)


@pytest.fixture(autouse=True)
def _isolate(monkeypatch, tmp_path):
    monkeypatch.setattr(ml_forecast, "MODELS_DIR", str(tmp_path))  # jangan baca model asli
    monkeypatch.delenv("REPORT_AI_ENABLED", raising=False)


def _rows(target_id, start_utc, end_utc, step_min=30, down_window=None, anomaly_every=None, rt=None):
    rows, t, i = [], start_utc, 0
    while t < end_utc:
        local = (t + timedelta(hours=7))
        is_down = bool(down_window and down_window[0] <= t < down_window[1])
        rows.append(CheckResult(
            target_id=target_id, checked_at=t, is_up=not is_down, status_code=None if is_down else 200,
            response_time_ms=None if is_down else (rt if rt is not None else 300.0 + local.hour * 5),
            is_anomaly=bool(anomaly_every and i % anomaly_every == 0 and not is_down),
        ))
        t += timedelta(minutes=step_min)
        i += 1
    return rows


async def _seed_august(db, target):
    down = (AUG_START + timedelta(days=9, hours=5), AUG_START + timedelta(days=9, hours=8))  # 3 jam
    db.add_all(_rows(target.id, AUG_START, AUG_END, down_window=down, anomaly_every=74))
    db.add_all(_rows(target.id, JUL_START, JUL_START + timedelta(days=2), rt=200.0))       # bulan sebelumnya
    db.add_all(_rows(target.id, JUL_END - timedelta(hours=5), JUL_END, rt=999.0))             # tepat sebelum batas
    db.add_all(_rows(target.id, AUG_END, AUG_END + timedelta(hours=5), rt=999.0))             # tepat setelah batas
    await db.commit()


def test_period_helpers_and_timezone_bounds():
    assert previous_period(datetime(2026, 9, 21, 10, 0, tzinfo=TZ)) == "2026-08"
    assert previous_period(datetime(2026, 1, 1, 9, 0, tzinfo=TZ)) == "2025-12"
    assert AUG_START == datetime(2026, 7, 31, 17, 0)   # 1 Agustus 00:00 WIB
    assert AUG_END == datetime(2026, 8, 31, 17, 0)     # 1 September 00:00 WIB
    assert period_label("2026-08") == "Agustus 2026"


async def test_stats_are_exact_and_respect_period_boundaries(db_session, sample_target):
    await _seed_august(db_session, sample_target)

    s = await collect_target_stats(db_session, sample_target, AUG_START, AUG_END, JUL_START, JUL_END, TZ)

    assert s["total_checks"] == 1488
    assert s["up_checks"] == 1482
    assert s["uptime_pct"] == pytest.approx(1482 / 1488 * 100)
    assert s["incident_count"] == 1
    assert s["downtime_minutes"] == pytest.approx(180)
    assert s["longest_incident_minutes"] == pytest.approx(180)
    assert s["anomaly_count"] > 0
    assert s["busiest_hour"] == 23
    assert s["max_rt"] < 999          # baris di luar periode tidak ikut
    assert s["prev_avg_rt"] is not None and s["prev_avg_rt"] < 999
    assert len(s["daily"]) == 31
    assert s["daily"][0]["date"] == "2026-08-01"
    assert len(s["incidents"]) == 1


async def test_incident_that_lasts_until_period_end_is_bounded(db_session, sample_target):
    db_session.add_all(_rows(sample_target.id, AUG_START, AUG_END, down_window=(AUG_END - timedelta(hours=2), AUG_END)))
    await db_session.commit()

    s = await collect_target_stats(db_session, sample_target, AUG_START, AUG_END, JUL_START, JUL_END, TZ)

    assert s["incident_count"] == 1
    assert s["downtime_minutes"] == pytest.approx(120)


async def test_target_without_data_yields_empty_stats(db_session, sample_target):
    s = await collect_target_stats(db_session, sample_target, AUG_START, AUG_END, JUL_START, JUL_END, TZ)
    assert s["total_checks"] == 0 and s["uptime_pct"] is None and s["daily"] == [] and s["incident_count"] == 0


async def test_charts_are_valid_png_and_capped(db_session, sample_target):
    await _seed_august(db_session, sample_target)
    s = await collect_target_stats(db_session, sample_target, AUG_START, AUG_END, JUL_START, JUL_END, TZ)

    without_forecast = render_charts("Agustus 2026", [s], TZ)
    s["forecast_series"] = [((AUG_END + timedelta(hours=i)).isoformat(), 300.0 + i) for i in range(168)]
    with_forecast = render_charts("Agustus 2026", [s], TZ)

    assert len(without_forecast) == 3 and len(with_forecast) == 4   # fleet + respons + uptime (+ prediksi)
    assert all(c["bytes"][:8] == b"\x89PNG\r\n\x1a\n" for c in with_forecast)
    assert all(len(c["caption"]) <= 1024 for c in with_forecast)

    many = [dict(s, name=f"T{i}", uptime_pct=99.0 + i / 100) for i in range(9)]
    capped = render_charts("Agustus 2026", many, TZ)
    assert len(capped) == 1 + monthly_report.MAX_CHART_TARGETS * 3


async def test_workbook_has_sheets_and_neutralises_formula_names(db_session, sample_target):
    sample_target.name = "=EVIL()"
    await _seed_august(db_session, sample_target)
    s = await collect_target_stats(db_session, sample_target, AUG_START, AUG_END, JUL_START, JUL_END, TZ)

    wb = load_workbook(io.BytesIO(build_report_workbook("Agustus 2026", [s])))

    assert wb.sheetnames == ["Ringkasan", "Harian", "Insiden"]
    assert wb["Ringkasan"].max_row == 2 and wb["Harian"].max_row == 32 and wb["Insiden"].max_row == 2
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                assert cell.data_type != "f"


# ---------- pengiriman ----------

@pytest.fixture
def sent(monkeypatch):
    calls = {"text": AsyncMock(return_value=True), "photos": AsyncMock(return_value=99), "doc": AsyncMock(return_value=True)}
    monkeypatch.setattr(monthly_report, "send_alert", calls["text"])
    monkeypatch.setattr(monthly_report, "send_photos", calls["photos"])
    monkeypatch.setattr(monthly_report, "send_document", calls["doc"])
    return calls


NOW = datetime(2026, 9, 1, 2, 0)  # 09:00 WIB tanggal 1


async def _link(db, user):
    user.telegram_chat_id = "chat-1"
    await db.commit()


async def _logs(db):
    return (await db.execute(select(MonthlyReportLog))).scalars().all()


async def test_report_is_delivered_once_then_idempotent(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)

    first = await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)
    second = await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)

    assert first.status == "sent" and second.status == "already_sent"
    assert sent["text"].call_count == 1 and sent["photos"].call_count == 1 and sent["doc"].call_count == 1
    text = sent["text"].call_args[0][1]
    assert "LAPORAN BULANAN" in text and "Agustus 2026" in text
    logs = await _logs(db_session)
    assert [(l.period, l.status) for l in logs] == [("2026-08", "sent")]


async def test_force_resends_without_duplicating_log(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)
    await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)

    again = await generate_and_send_report(db_session, sample_user, mode="previous", force=True, now=NOW)

    assert again.status == "sent" and sent["text"].call_count == 2
    assert len(await _logs(db_session)) == 1


async def test_current_month_mode_reports_month_to_date_and_is_not_logged(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    start = period_bounds_utc("2026-09", TZ)[0]
    db_session.add_all(_rows(sample_target.id, start, start + timedelta(days=3)))
    await db_session.commit()

    result = await generate_and_send_report(db_session, sample_user, mode="current", now=start + timedelta(days=3))

    assert result.status == "sent"
    assert "September 2026" in sent["text"].call_args[0][1]
    assert await _logs(db_session) == []


async def test_not_linked_inactive_no_targets_and_no_data(db_session, sample_user, sent):
    assert (await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)).status == "not_linked"

    await _link(db_session, sample_user)
    assert (await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)).status == "no_targets"

    target = MonitorTarget(user_id=sample_user.id, name="Kosong", url="https://x.example.com")
    db_session.add(target)
    await db_session.commit()
    assert (await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)).status == "no_data"
    assert [(l.status) for l in await _logs(db_session)] == ["skipped"]

    sample_user.is_active = False
    await db_session.commit()
    assert (await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)).status == "inactive"
    sent["text"].assert_not_called()


async def test_failed_delivery_is_retried_and_log_row_is_updated(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)
    sent["text"].return_value = False

    failed = await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)
    assert failed.status == "failed"
    assert [l.status for l in await _logs(db_session)] == ["failed"]

    sent["text"].return_value = True
    retried = await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)
    assert retried.status == "sent"
    assert [l.status for l in await _logs(db_session)] == ["sent"]


async def test_partial_failure_of_photos_or_excel_still_counts_as_sent_with_detail(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)
    sent["photos"].return_value = 0
    sent["doc"].return_value = False

    result = await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)

    assert result.status == "sent"
    assert "gagal" in result.detail


async def test_report_when_delivery_functions_raise_does_not_propagate(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)
    sent["photos"].side_effect = RuntimeError("boom")

    result = await generate_and_send_report(db_session, sample_user, mode="previous", now=NOW)

    assert result.status == "sent"


# ---------- penjadwalan / catch-up ----------

class _SessionFactory:
    def __init__(self, session):
        self._s = session

    def __call__(self):
        return self

    async def __aenter__(self):
        return self._s

    async def __aexit__(self, *a):
        return False


async def test_scheduler_waits_until_send_hour_on_the_first(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)

    early = await ensure_monthly_reports(now=datetime(2026, 9, 1, 0, 30), session_factory=_SessionFactory(db_session))  # 07:30 WIB
    assert early == [] and sent["text"].call_count == 0

    on_time = await ensure_monthly_reports(now=NOW, session_factory=_SessionFactory(db_session))
    assert on_time == [(sample_user.id, "sent")] and sent["text"].call_count == 1


async def test_catch_up_sends_missed_report_days_later_but_only_once(db_session, sample_user, sample_target, sent):
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)
    later = datetime(2026, 9, 15, 5, 0)  # server baru menyala tanggal 15

    first = await ensure_monthly_reports(now=later, session_factory=_SessionFactory(db_session))
    second = await ensure_monthly_reports(now=later + timedelta(hours=1), session_factory=_SessionFactory(db_session))

    assert first == [(sample_user.id, "sent")]
    assert second == [(sample_user.id, "already_sent")]
    assert sent["text"].call_count == 1


async def test_one_broken_user_does_not_block_the_others(db_session, sample_user, sample_target, sent, monkeypatch):
    other = User(email="o@example.com", hashed_password="x", telegram_chat_id="chat-2")
    db_session.add(other)
    await db_session.commit()
    await _link(db_session, sample_user)
    await _seed_august(db_session, sample_target)

    original = monthly_report.generate_and_send_report

    async def flaky(db, user, **kw):
        if user.id == sample_user.id:
            raise RuntimeError("crash")
        return await original(db, user, **kw)

    monkeypatch.setattr(monthly_report, "generate_and_send_report", flaky)

    results = await ensure_monthly_reports(now=NOW, session_factory=_SessionFactory(db_session))

    assert dict(results)[sample_user.id] == "error"
    assert dict(results)[other.id] in ("no_targets", "no_data")


def test_schedule_registers_hourly_job():
    from apscheduler.schedulers.asyncio import AsyncIOScheduler
    scheduler = AsyncIOScheduler()
    monthly_report.schedule_report_jobs(scheduler)
    assert {j.id for j in scheduler.get_jobs()} == {"monthly_reports", "monthly_reports_startup"}


# ---------- pengiriman Telegram (foto/dokumen) ----------

class _FakeBot:
    calls = []
    fail = False

    def __init__(self, token=None, **kw):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def send_media_group(self, chat_id, media, **kw):
        if _FakeBot.fail:
            raise RuntimeError("x")
        _FakeBot.calls.append(("group", len(media)))

    async def send_photo(self, chat_id, photo, caption=None, **kw):
        if _FakeBot.fail:
            raise RuntimeError("x")
        _FakeBot.calls.append(("photo", 1))

    async def send_document(self, chat_id, document, filename=None, caption=None, **kw):
        if _FakeBot.fail:
            raise RuntimeError("x")
        _FakeBot.calls.append(("doc", filename))


@pytest.fixture
def fake_bot(monkeypatch):
    _FakeBot.calls, _FakeBot.fail = [], False
    monkeypatch.setattr(telegram_bot, "Bot", _FakeBot)
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "1" * 9 + ":" + "A" * 30)
    return _FakeBot


async def test_send_photos_chunks_into_groups_of_ten_and_single_uses_send_photo(fake_bot):
    photos = [(b"x", f"c{i}") for i in range(11)]
    assert await telegram_bot.send_photos("c", photos) == 11
    assert fake_bot.calls == [("group", 10), ("photo", 1)]


async def test_send_photos_counts_only_delivered_and_never_raises(fake_bot):
    fake_bot.fail = True
    assert await telegram_bot.send_photos("c", [(b"x", "a"), (b"y", "b")]) == 0


async def test_send_document_and_no_token(fake_bot, monkeypatch):
    assert await telegram_bot.send_document("c", b"data", "laporan.xlsx", "cap") is True
    assert fake_bot.calls[-1] == ("doc", "laporan.xlsx")
    monkeypatch.delenv("TELEGRAM_BOT_TOKEN")
    assert await telegram_bot.send_document("c", b"data", "laporan.xlsx", "cap") is False
    assert await telegram_bot.send_photos("c", [(b"x", "a")]) == 0
