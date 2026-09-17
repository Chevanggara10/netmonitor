from openpyxl import load_workbook
import io

from app.excel_export import export_checks_to_excel
from tests.conftest import add_check_results


async def test_export_with_data_has_correct_columns_and_rows(db_session, sample_target):
    await add_check_results(db_session, sample_target.id, count=5)

    content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=30)

    wb = load_workbook(io.BytesIO(content))
    ws = wb["Raw Checks"]
    header = [cell.value for cell in ws[1]]
    assert header == [
        "target_name", "checked_at", "status_code", "response_time_ms",
        "is_up", "error_message", "is_anomaly", "anomaly_z_score",
    ]
    # 1 header row + 5 data rows
    assert ws.max_row == 6
    assert ws.cell(row=2, column=1).value == "Test Target"


async def test_export_with_no_data_has_header_only(db_session, sample_target):
    content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=30)

    wb = load_workbook(io.BytesIO(content))
    ws = wb["Raw Checks"]
    assert ws.max_row == 1  # cuma header, tidak crash


async def test_export_has_hourly_aggregate_sheet(db_session, sample_target):
    await add_check_results(db_session, sample_target.id, count=3)

    content = await export_checks_to_excel(sample_target.id, sample_target.name, db_session, days=30)

    wb = load_workbook(io.BytesIO(content))
    assert "Hourly Aggregate" in wb.sheetnames
    ws = wb["Hourly Aggregate"]
    header = [cell.value for cell in ws[1]]
    assert header == ["hour", "avg_response_time_ms", "total_checks", "uptime_percent"]
