"""
Script training offline: baca Excel hasil export (app/excel_export.py),
resample response_time_ms ke per-jam, fit model Holt-Winters
(statsmodels ExponentialSmoothing), simpan ke training/models/{target_id}.pkl.

Dijalankan manual atau lewat scheduler OS (cron/Task Scheduler), TIDAK
berjalan di dalam proses web server -- training bisa makan waktu lebih
lama dari 1x siklus check biasa.

Pengaman data (agar model tidak belajar dari data palsu/rusak):
- Celah data panjang (mis. laptop mati berhari-hari) TIDAK diinterpolasi;
  hanya segmen kontigu TERBARU yang dipakai.
- Outlier ekstrem (timeout puluhan detik) dipotong dengan pagar Tukey.
- Tren diredam (damped) supaya prediksi tidak meledak linear tanpa batas.
- Hasil fit divalidasi (tidak boleh NaN/inf) dan disimpan ATOMIK, jadi
  model lama tetap utuh kalau training gagal / web sedang membaca file.

Cara pakai:
    python training/train_forecast.py --excel laporan.xlsx --target-id 1
"""
import argparse
import os
import sys
import tempfile
import time
from datetime import datetime

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from statsmodels.tsa.holtwinters import ExponentialSmoothing

# Minimum titik data (per jam) sebelum training dianggap layak dipercaya.
# Di bawah ini, hasil fit terlalu tidak reliable untuk dipakai forecast.
MIN_DATA_POINTS = 48
# Celah <= ini (jam) diinterpolasi; lebih panjang = data dianggap terputus.
MAX_INTERPOLATE_GAP_HOURS = 6
# Format bundle .pkl (dibaca app/ml_forecast.py).
MODEL_FORMAT_VERSION = 1
REQUIRED_COLUMNS = ("checked_at", "response_time_ms")
# Absolut & sama dengan app.ml_forecast.MODELS_DIR, apa pun folder tempat script dijalankan.
DEFAULT_MODELS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "models")


class TrainingDataError(ValueError):
    """Data/hasil training tidak layak; pesan ditujukan untuk dibaca operator."""


def read_raw_checks(excel_path: str) -> list[dict]:
    try:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as exc:
        raise TrainingDataError(f"File Excel tidak bisa dibuka: {exc}") from exc

    try:
        if "Raw Checks" not in wb.sheetnames:
            raise TrainingDataError(
                f'Sheet "Raw Checks" tidak ditemukan (sheet yang ada: {wb.sheetnames}). '
                "Pakai file hasil export dari endpoint /export/excel."
            )
        rows = list(wb["Raw Checks"].iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise TrainingDataError(f"Kolom wajib tidak ada di sheet Raw Checks: {', '.join(missing)}")
    return [dict(zip(header, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def _empty_series() -> pd.Series:
    return pd.Series(dtype=float, index=pd.DatetimeIndex([]))


def resample_to_hourly(raw_rows: list[dict]) -> pd.Series:
    """
    raw_rows: list of dict dengan minimal key 'checked_at' (ISO string atau
    datetime bawaan Excel) dan 'response_time_ms' (float; None untuk baris down).
    Return: rata-rata response_time_ms per jam. Jam tanpa data = NaN
    (TIDAK diinterpolasi di sini -- lihat prepare_training_series).
    """
    if not raw_rows:
        return _empty_series()

    df = pd.DataFrame(raw_rows)
    df["checked_at"] = pd.to_datetime(df["checked_at"], errors="coerce")
    df["response_time_ms"] = pd.to_numeric(df["response_time_ms"], errors="coerce")
    df = df.dropna(subset=["checked_at"])
    if df.empty:
        return _empty_series()

    if df["checked_at"].dt.tz is not None:
        df["checked_at"] = df["checked_at"].dt.tz_convert("UTC").dt.tz_localize(None)

    df = df.set_index("checked_at").sort_index()
    return df["response_time_ms"].resample("h").mean()


def prepare_training_series(raw_rows: list[dict]) -> pd.Series:
    """
    Deret per jam siap latih: buang NaN di ujung, potong pada celah panjang
    (ambil segmen terbaru saja), interpolasi celah pendek.
    """
    hourly = resample_to_hourly(raw_rows)
    if hourly.empty or not hourly.notna().any():
        return _empty_series()

    hourly = hourly.loc[hourly.first_valid_index():hourly.last_valid_index()]

    isna = hourly.isna()
    run_id = (isna.astype(int).diff().fillna(0) != 0).cumsum()
    run_len = isna.astype(int).groupby(run_id).transform("sum")
    long_gap = isna & (run_len > MAX_INTERPOLATE_GAP_HOURS)
    if long_gap.any():
        last_gap_pos = int(np.flatnonzero(long_gap.to_numpy())[-1])
        hourly = hourly.iloc[last_gap_pos + 1:]

    return hourly.interpolate(method="linear", limit_area="inside").dropna()


def clip_outliers(series: pd.Series) -> pd.Series:
    """Pagar Tukey ekstrem (Q3 + 3*IQR): membuang lonjakan timeout, menjaga puncak wajar."""
    if series.empty:
        return series
    q1, q3 = series.quantile(0.25), series.quantile(0.75)
    return series.clip(lower=0, upper=q3 + 3 * (q3 - q1))


def _fit_model(series: pd.Series):
    return ExponentialSmoothing(
        series, trend="add", damped_trend=True, seasonal="add", seasonal_periods=24,
    ).fit()


def _atomic_dump(bundle: dict, final_path: str) -> None:
    directory = os.path.dirname(final_path) or "."
    fd, tmp_path = tempfile.mkstemp(dir=directory, prefix=".model-", suffix=".tmp")
    os.close(fd)
    try:
        joblib.dump(bundle, tmp_path)
        for attempt in range(5):
            try:
                os.replace(tmp_path, final_path)
                return
            except PermissionError:
                # Windows: web server sedang membuka file lama sesaat.
                if attempt == 4:
                    raise
                time.sleep(0.3)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def train_and_save(excel_path: str, target_id: int, output_dir: str = DEFAULT_MODELS_DIR) -> str | None:
    raw_rows = read_raw_checks(excel_path)
    series = prepare_training_series(raw_rows)

    if len(series) < MIN_DATA_POINTS:
        print(f"[skip] Target {target_id}: cuma {len(series)} titik data per jam yang kontigu "
              f"(celah > {MAX_INTERPOLATE_GAP_HOURS} jam memutus data), "
              f"butuh minimal {MIN_DATA_POINTS}. Tidak training.")
        return None

    series = clip_outliers(series).asfreq("h")

    try:
        model = _fit_model(series)
    except Exception as exc:
        raise TrainingDataError(f"Gagal melatih model: {exc}") from exc

    if not np.isfinite(np.asarray(model.fittedvalues, dtype=float)).all():
        raise TrainingDataError(
            "Model tidak stabil (hasil fit mengandung NaN/inf); model lama dipertahankan."
        )

    bundle = {
        "version": MODEL_FORMAT_VERSION,
        "model": model,
        "last_data_at": series.index[-1].isoformat(),
        "trained_at": datetime.utcnow().isoformat(),
        "n_points": int(len(series)),
    }

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f"{target_id}.pkl")
    _atomic_dump(bundle, output_path)
    print(f"[ok] Model target {target_id} disimpan ke {output_path} ({len(series)} titik data).")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Training model forecast dari Excel export Netmonitor.")
    parser.add_argument("--excel", required=True, help="Path ke file .xlsx hasil export")
    parser.add_argument("--target-id", type=int, required=True, help="ID target monitoring")
    parser.add_argument("--output-dir", default=DEFAULT_MODELS_DIR)
    args = parser.parse_args()

    try:
        train_and_save(args.excel, args.target_id, args.output_dir)
    except TrainingDataError as exc:
        print(f"[error] {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
